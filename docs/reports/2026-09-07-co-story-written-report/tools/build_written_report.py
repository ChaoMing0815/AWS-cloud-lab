#!/usr/bin/env python3
"""Build the Co-Story final written report DOCX from the Markdown manuscript."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BODY_FONT = "Arial Unicode MS"
HEADING_FONT = "Arial Unicode MS"
MONO_FONT = "Liberation Mono"


def set_run_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_east_asian_paragraph_rules(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for tag, value in (("kinsoku", "1"), ("overflowPunct", "0")):
        node = ppr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            ppr.append(node)
        node.set(qn("w:val"), value)


def add_field(paragraph, instruction: str, display: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate))
    if display:
        text = OxmlElement("w:t")
        text.text = display
        run._r.append(text)
    run._r.append(end)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = HEADING_FONT
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True

    heading_specs = {
        "Heading 1": (18, 18, 10),
        "Heading 2": (15, 14, 8),
        "Heading 3": (13, 10, 6),
    }
    for style_name, (size, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True
        style.paragraph_format.page_break_before = False

    if "Figure Caption" not in styles:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = BODY_FONT
    cap.font.size = Pt(10)
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.widow_control = True

    if "Code Inline" not in styles:
        code = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        code.font.name = MONO_FONT
        code.font.size = Pt(10.5)
        code.font.color.rgb = RGBColor(32, 32, 32)

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        add_field(p, "PAGE", "1")
        for run in p.runs:
            set_run_font(run, BODY_FONT, 9)


def add_inline_runs(paragraph, text: str, *, bold=False, size=12) -> None:
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.style = "Code Inline"
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, BODY_FONT, size, bold=bold)


def image_size(path: Path, max_width_in: float = 6.0, max_height_in: float = 6.2):
    with Image.open(path) as im:
        width, height = im.size
    ratio = width / height
    if ratio < 1:
        max_height_in = min(max_height_in, 5.5)
    width_in = max_width_in
    height_in = width_in / ratio
    if height_in > max_height_in:
        height_in = max_height_in
        width_in = height_in * ratio
    return Inches(width_in), Inches(height_in)


def add_image(doc: Document, path: Path, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    width, height = image_size(path)
    run = p.add_run()
    shape = run.add_picture(str(path), width=width, height=height)
    shape._inline.docPr.set("descr", alt)


def add_cover_and_toc(doc: Document, title_text: str, subtitle: str) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(74)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    run = title.add_run(title_text)
    set_run_font(run, HEADING_FONT, 26, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Pt(0)
    sub.paragraph_format.space_after = Pt(36)
    run = sub.add_run(subtitle)
    set_run_font(run, HEADING_FONT, 15)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.first_line_indent = Pt(0)
    meta.paragraph_format.line_spacing = 1.5
    add_inline_runs(meta, "Release v1.1.6\n2026 年 9 月 8 日", size=12)

    doc.add_page_break()
    toc_title = doc.add_paragraph("目錄", style="Heading 1")
    toc_title.paragraph_format.first_line_indent = Pt(0)
    toc_entries = (
        "摘要",
        "第一章　緒論",
        "第二章　需求分析與系統設計",
        "第三章　AWS 架構設計",
        "第四章　系統實作",
        "第五章　測試、部署與維運",
        "第六章　成果、限制與未來方向",
        "參考資料",
        "附錄",
    )
    for entry in toc_entries:
        toc = doc.add_paragraph()
        toc.paragraph_format.first_line_indent = Pt(0)
        toc.paragraph_format.space_after = Pt(7)
        add_inline_runs(toc, entry, size=12)
    doc.add_page_break()


def build(manuscript: Path, output: Path) -> None:
    lines = manuscript.read_text(encoding="utf-8").splitlines()
    if len(lines) < 3 or not lines[0].startswith("# "):
        raise ValueError("Unexpected manuscript title block")

    doc = Document()
    style_document(doc)
    add_cover_and_toc(doc, lines[0][2:].strip(), lines[2].strip())

    base = manuscript.parent
    previous_was_image = False
    for raw in lines[3:]:
        line = raw.strip()
        if not line:
            previous_was_image = False
            continue

        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, relative = image_match.groups()
            image_path = (base / relative).resolve()
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            add_image(doc, image_path, alt)
            previous_was_image = True
            continue

        if line.startswith("圖 "):
            p = doc.add_paragraph(style="Figure Caption")
            add_inline_runs(p, line, size=10)
            previous_was_image = False
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline_runs(p, line[2:].strip(), bold=True, size=18)
            set_east_asian_paragraph_rules(p)
            continue

        if line.startswith("## "):
            text = line[3:].strip()
            style = "Heading 1" if text == "摘要" else "Heading 2"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline_runs(p, text, bold=True, size=18 if style == "Heading 1" else 15)
            set_east_asian_paragraph_rules(p)
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline_runs(p, line[4:].strip(), bold=True, size=13)
            set_east_asian_paragraph_rules(p)
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.widow_control = True
            add_inline_runs(p, numbered.group(2))
            set_east_asian_paragraph_rules(p)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        add_inline_runs(p, line)
        set_east_asian_paragraph_rules(p)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = "共演計劃多人 AI 協作故事遊戲 AWS 雲端系統實作"
    doc.core_properties.subject = "AWS 雲端工程師培訓期末專題報告"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_written_report.py MANUSCRIPT OUTPUT")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
